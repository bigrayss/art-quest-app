#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""任务总清单（xlsx）：首访库 + 主库，每道题一行——题目、六个年龄段打钩、参考图缩略图、维度考察、来源依据。

    .venv/bin/python tools/gen_task_catalog.py        # 写 docs/任务总清单.xlsx

内容全部从代码里抽（题面：missions_v2 / missions_bank；依据：artquest/taskspace），改了题面或来源重跑一次就对上。
年龄段的钩是**默认值**（规则见「说明」页），可以直接在表里改。
"""
import io
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from artquest.missions_bank import BANK_VERSION  # noqa: E402
from artquest.quests import QUESTS  # noqa: E402
from artquest.reconstruct import draw_stimulus  # noqa: E402
from artquest.scoring import DIMENSIONS  # noqa: E402
from artquest.taskspace.blueprints import FLOW_SOURCES  # noqa: E402
from artquest.taskspace.space import initialize  # noqa: E402
from artquest.taskspace.task_sources import CATALOG_BIB, FAMILY_NOTES, TASK_DESIGN  # noqa: E402

OUT = ROOT / "docs" / "任务总清单.xlsx"
STATIC = ROOT / "static"
DIM_ZH = {d["key"]: d["zh"] for d in DIMENSIONS}
BANDS = [("B1", "6–7 岁"), ("B2", "8–10 岁"), ("B3", "11–12 岁"), ("B4", "13–14 岁"), ("B5", "15–16 岁"), ("B6", "17 岁及以上")]
FAMILY_TITLE = {
    "M0": "基线", "M1": "博物馆修复师", "M2": "探险家速写", "M3": "失落的碎片", "M4": "变异物体实验室",
    "M5": "融合发明家", "M6": "情绪世界", "M7": "线条冒险", "M8": "不可能世界", "M9": "故事挑战",
    "OB": "观察与表现", "ST": "个人表达与叙事", "FG": "图形片段创作", "TR": "物品创意改造", "DS": "视觉传达设计",
}
SOURCE_SHORT = {"kb2022": "课标 2022", "kidsartbench": "KidsArtBench", "pisa2022-ct-framework": "PISA 2022 框架",
                "pisa2022-vol3": "PISA 2022 Vol. III", "urban2005": "Urban 2005 TCT-DP", "ttct": "TTCT 手册 2018"}
THUMB_W = 160          # 缩略图宽（像素）
NOTES = [
    "由 tools/gen_task_catalog.py 生成：主库 TASK_VERSION {tv}；首访库 BANK_VERSION {bv}。正文（设计思路与来源、家族依据、来源文献、本页）是 2026-10-12 用户定稿的文字，10-13 回写进 artquest/taskspace/task_sources.py 作为数据源；改文字改那里再重跑，别手改 xlsx。不改运行题库、程序路由或评分配置。",
    "",
    "阅读顺序与修改范围：先看任务清单，再按片段ID查来源片段／来源文献；家族依据说明共性。本次保留30道首访题＋80道主库题，以及原题面、年龄标记、图片和维度配置。",
    "✓ = 原表的课标学段关联标记，不是经验证的适龄结论。保留原映射：1—2年级→B1/B2；3—5年级→B2/B3；6—7年级→B3/B4；8—9年级→B4/B5。以上映射是项目索引规则，不是课标官方年龄换算。",
    "○ = 原表项目拟适用范围。首访按所属段，主库小学版B1—B3、初中版B4—B6、指定固定参照按原标记。本次均未重算；这些标记不是作品水平或实际首访路由配置。",
    "B6为17岁及以上，包含17岁未成年人，不能统称成人组。义务教育课标不能直接证明B6适宜性，也不能提供成人常模；年龄不等于绘画经验或能力。",
    "",
    "“设计思路与来源”分三层：设计意图（本项目希望做什么）；来源对应（释义，附可定位片段ID）；项目改编与边界（新增条件、迁移范围与不能预先得出的结论）。",
    "来源对应主要保留原表的内容关联，并为图形片段活动明确Urban范式出处。正文引用是释义，不冒充逐字引文；原始转录与节录仍在“来源片段”页，未作静默改写。",
    "家族依据不替代实际题面。来源支持学习方向，不自动规定每个对象、配色、部件对应或难度。“想想”“可以”“尽量”等保留原题强度，不由设计说明加成隐藏要求。",
    "原项目记录位置：docs/任务空间/README.md；artquest/taskspace/task_sources.py（TASK_DESIGN / FAMILY_NOTES / CATALOG_BIB）、blueprints.py、sources.py。改动已随脚本长期保留。",
    "保留原表形成顺序：任务写于2026-10-01，来源后来回溯匹配；此处说明“哪些内容与已有资料对应”，不改写为“每道题都预先从资料演绎生成”。原记录见 docs/来源核查_2026-10-12.md。",
    "来源片段共44条，沿用原表转录/节录。本次核对相关课程页及关键方法范围；不宣称44条重新逐字校验。页码按各文献定位，课标为正文页＝PDF页序−7。",
    "",
    "图片保留原表58张及其对应行：主库M1/M2/M4/M6来自static/refs/，M3为原表矢量规则图，首访来自static/refs/bank/。本次不换图、不裁切，也不调整题面中的对象关系。",
    "",
    "流程边界：本项目“首稿→反馈→修改自己的作品”属于学习支持设计；PISA可提供想法改进的概念参照，但原测评主要改进他人给定方案。单幅产出是本项目实现选择，不由PISA单创意要求直接推出。",
    "评价边界：KidsArtBench评作品；任务与九维的关联不等于分数、权重或能力结论。本次保留原主要/次要等配置，不实施降级或改评分契约；原表术语需要进一步统一时应另作版本变更。",
    "研究参照：主库不是备用池。指定固定任务可承担共同参照；“固定”“锚定”“同家族”均不意味着难度等值、跨年龄可比或测量有效，这些用途需要实际设计与数据支持。",
    "本次文字处理：将“最难、最低/最高阶”等无数据断言改成具体条件；将部件映射改成开放路径；区分课程对应、范式借鉴和项目自定。内容一致性得到澄清，不预填教师审核或学生验证结果。"
]
TICK, PROPOSED = "✓", "○"     # ✓ 课标片段推出来的学段；○ 项目拟适用（来源推不出来）


def cite(ex_id, excerpts):
    e = excerpts.get(ex_id)
    if not e:
        return ex_id or ""
    src = SOURCE_SHORT.get(e.get("source"), e.get("source"))
    page = f" p.{e['page_body']}" if e.get("page_body") else ""
    head = f" {e['heading']}" if e.get("heading") else ""
    return f"{src}{page}{head}"


# 课标学段 → 年龄段：一(1–2 年级)≈6–8 岁、二(3–5)≈8–11、三(6–7)≈11–13、四(8–9)≈13–15。课标到 9 年级为止，B6 永远推不出来。
STAGE_BANDS = {"一": {"B1", "B2"}, "二": {"B2", "B3"}, "三": {"B3", "B4"}, "四": {"B4", "B5"}}


def bands_from_sources(c, excerpts):
    """✓ 的来源：这道题**自己**（scope=task）引的、关系为「原文明确规定」、人判为「支持」的课标片段属于哪个学段。
    家族共用的句子不算——不然同家族 8 道题的 ✓ 一模一样，又回到按家族挂。"""
    out = set()
    for r in c.get("source_relations") or []:
        if r.get("scope") != "task" or r.get("relation") != "原文明确规定" or (r.get("support") or "支持") != "支持":
            continue
        e = excerpts.get(r.get("excerpt") or "")
        if e and e.get("source") == "kb2022":
            out |= STAGE_BANDS.get(e.get("band"), set())
    return out


def bands_proposed(c):
    """○ 的来源：项目自己定的适用范围——首访题是它所属的那一段；主库小学版 B1–B3、初中版 B4–B6、锚定题六段。"""
    if c.get("bank"):
        return {c.get("band")}
    tiers = set(c.get("tiers") or [])
    if c.get("anchor") or tiers == {"simple", "full"}:
        return {b for b, _ in BANDS}
    out = set()
    if "simple" in tiers:
        out |= {"B1", "B2", "B3"}
    if "full" in tiers:
        out |= {"B4", "B5", "B6"}
    return out


def thumb(c):
    """参考图缩略图：文件图直接缩；碎片题按 app 同一套矢量规则画出来再缩。没有参考图返回 None。"""
    m = c.get("materials") or {}
    q = next((x for x in QUESTS if x["id"] == c["task_id"]), {})
    stim = q.get("stimulus") or {}
    if m.get("kind") == "fragments":
        img = Image.new("RGB", (1024, 704), "white")
        draw_stimulus(img, stim)
    elif m.get("kind") == "reference" and stim.get("file"):
        p = STATIC / stim["file"].split("/static/", 1)[-1]
        if not p.exists():
            return None
        img = Image.open(p).convert("RGB")
    else:
        return None
    h = max(1, round(img.height * THUMB_W / img.width))
    img = img.resize((THUMB_W, h), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf, h


def material_text(c):
    m = c.get("materials") or {}
    kind = m.get("kind")
    if kind in (None, "none"):
        return "无参考图（空白画布）"
    spec = m.get("spec") or {}
    q = next((x for x in QUESTS if x["id"] == c["task_id"]), {})
    stim = q.get("stimulus") or {}
    if kind == "fragments":
        return f"碎片印在初始画布上（{'、'.join(spec.get('objects') or [])}；擦不掉）"
    mode = {"always": "常显", "on_demand": "点开看"}.get(m.get("mode"), "")
    parts = [f"{stim.get('file', '')}（{mode}）"]
    if spec.get("objects"):
        parts.append("画面：" + "、".join(spec["objects"]))
    if spec.get("status"):
        parts.append(spec["status"])
    return "\n".join(parts)


def dims_text(c):
    ds = c.get("dims") or []
    p = [DIM_ZH.get(d["dim"], d["dim"]) for d in ds if d.get("relation") == "primary"]
    s = [DIM_ZH.get(d["dim"], d["dim"]) for d in ds if d.get("relation") == "secondary"]
    out = []
    if p:
        out.append("主要：" + "、".join(p))
    if s:
        out.append("次要：" + "、".join(s))
    return "\n".join(out) if out else "（探索性，不作主要考察）"


def sources_text(c, excerpts):
    """清单那一列的正文就是 TASK_DESIGN（设计意图 / 来源对应 / 项目改编与边界），来源已经写在里面，不再另拼。"""
    return TASK_DESIGN.get(c["task_id"], "")


def main():
    ts = initialize()
    excerpts = ts["excerpts"]
    cands = [c for c in ts["candidates"] if not c.get("draft")]
    fam_order = {"OB": 0, "ST": 1, "FG": 2, "TR": 3, "DS": 4}          # 首访每段五道的固定顺序
    bank = sorted((c for c in cands if c.get("bank")), key=lambda c: (c.get("band"), fam_order.get(c["family"], 9), c["task_id"]))
    main_lib = sorted((c for c in cands if not c.get("bank")), key=lambda c: c["task_id"])

    wb = Workbook()
    ws = wb.active
    ws.title = "任务清单"
    head = ["库", "编号", "题名", "题目"] + [f"{b} {age}" for b, age in BANDS] + ["参考图", "参考图说明", "维度考察", "设计思路与来源"]
    ws.append(head)
    widths = [10, 11, 14, 44, 8, 8, 8, 8, 8, 9, 24, 34, 18, 70]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    hfill = PatternFill("solid", fgColor="F2E9DC")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = hfill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "E2"

    wrap = Alignment(vertical="top", wrap_text=True)
    center = Alignment(horizontal="center", vertical="center")
    img_col = get_column_letter(head.index("参考图") + 1)
    row = 2
    for lib, rows in (("首访库", bank), ("主库", main_lib)):
        for c in rows:
            cp = c.get("copy") or {}
            fam = c["family"]
            lib_cell = f"{lib}\n{fam} {FAMILY_TITLE.get(fam, '')}"
            src_b, prop_b = bands_from_sources(c, excerpts), bands_proposed(c)
            title = cp.get("title", "")
            text = cp.get("instruction", "")
            if c.get("requirements"):
                text += "\n要做到：" + c["requirements"]
            values = [lib_cell, c["task_id"], title, text] + [(TICK if b in src_b else (PROPOSED if b in prop_b else "")) for b, _ in BANDS] \
                + ["", material_text(c), dims_text(c), sources_text(c, excerpts)]
            ws.append(values)
            for i, cell in enumerate(ws[row], 1):
                cell.alignment = center if 5 <= i <= 10 else wrap
            t = thumb(c)
            h_px = 48
            if t:
                buf, h = t
                xi = XLImage(buf)
                xi.width, xi.height = THUMB_W, h
                ws.add_image(xi, f"{img_col}{row}")
                h_px = max(h_px, h + 6)
            # 行高按内容估：缩略图高度和各列折行后的行数取大（中文一个字约占两个列宽单位）
            def wrapped(v, width):
                return sum(max(1, -(-len(line) * 2 // width)) for line in str(v or "").split("\n"))
            lines = max(wrapped(values[i], widths[i]) for i in (0, 3, 11, 12, 13))
            ws.row_dimensions[row].height = max(h_px * 0.75, min(409, 15 * lines + 6))
            row += 1

    # 家族依据：一家族一行，不再复制到每道题
    wsf = wb.create_sheet("家族依据")
    wsf.append(["家族", "蓝图目标", "不能改的条件", "来源对应与改编范围（释义）", "流程依据与解释边界"])
    for cell in wsf[1]:
        cell.font = Font(bold=True)
        cell.fill = hfill
    for w, col in zip([18, 40, 40, 70, 50], "ABCDE"):
        wsf.column_dimensions[col].width = w
    seen = []
    for c in bank + main_lib:
        if c["family"] in seen:
            continue
        seen.append(c["family"])
        bp = ts["blueprints"].get(c["blueprint"]) or {}
        fn = FAMILY_NOTES.get(c["family"], {})
        wsf.append([f"{c['family']} {FAMILY_TITLE.get(c['family'], '')}", bp.get("goal", ""), "；".join(bp.get("invariants") or []),
                    fn.get("sources", ""), fn.get("flow", "")])
    for r in wsf.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = wrap
    wsf.freeze_panes = "A2"

    # 来源片段
    ws2 = wb.create_sheet("来源片段")
    ws2.append(["片段 id", "来源", "正文页", "PDF 页", "标题", "原文"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = hfill
    for w, col in zip([24, 18, 8, 8, 34, 90], "ABCDEF"):
        ws2.column_dimensions[col].width = w
    for eid, e in excerpts.items():
        ws2.append([eid, SOURCE_SHORT.get(e.get("source"), e.get("source")), e.get("page_body"), e.get("page_pdf"),
                    e.get("heading"), (e.get("text") or "").strip()])
    for r in ws2.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = wrap
    ws2.freeze_panes = "A2"

    # 来源文献
    ws3 = wb.create_sheet("来源文献")
    ws3.append(["id", "文献", "作用与适用边界", "定位与核对范围", "链接"])
    for cell in ws3[1]:
        cell.font = Font(bold=True)
        cell.fill = hfill
    for w, col in zip([20, 50, 50, 50, 50], "ABCDE"):
        ws3.column_dimensions[col].width = w
    for b in CATALOG_BIB:
        ws3.append([b["id"], b["citation"], b["role"], b["locator"], b["url"]])
    for r in ws3.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = wrap

    # 说明（用户 2026-10-12 定稿的文字，见 task_sources.py 顶部的说明）
    ws4 = wb.create_sheet("说明")
    for n in NOTES:
        ws4.append([n.format(tv=ts["task_version"], bv=BANK_VERSION)])
    ws4.column_dimensions["A"].width = 120
    for r in ws4.iter_rows():
        for cell in r:
            cell.alignment = wrap

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"写到 {OUT.relative_to(ROOT)}：首访 {len(bank)} + 主库 {len(main_lib)} = {row - 2} 行；片段 {len(excerpts)} 条")


if __name__ == "__main__":
    main()
