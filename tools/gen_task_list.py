# -*- coding: utf-8 -*-
"""生成给专业老师看的任务清单：`docs/任务清单.docx` 和 `.pdf`。

内容从 `artquest/missions.py` 抽（那是任务原文唯一的家），所以老师改了意见、我改了代码、
重跑一次，文件就和 app 里孩子看到的一字不差。每个家族只说三样：研究目的、主要考察、任务本身。

    python3 tools/gen_task_list.py          # 需要 google-chrome（出 PDF）和 python-docx（出 DOCX）
"""
import html
import os
import subprocess
import sys
import tempfile
from collections import OrderedDict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.environ.setdefault("ARTQUEST_SCORER", "heuristic")
os.environ.setdefault("ARTQUEST_FEEDBACK", "template")

from artquest import missions as _m  # noqa: E402
from artquest.quests import QUESTS, families  # noqa: E402
from artquest.scoring import DIMENSIONS  # noqa: E402

ZH = {d["key"]: d["zh"] for d in DIMENSIONS}
STYLE_ZH = {"minimal": "极简说法", "story": "故事说法", "challenge": "挑战说法"}

# 研究目的，用老师听得懂的话说。代码里的 research_goal 是给研究员的缩写。
AIM = {
    "M0": "开放创作。不设任何限制，看孩子在完全自由时怎么选题、怎么组织画面。",
    "M1": "根据一张受损的参考图重建场景。看的是孩子如何组织整体结构、先抓大局还是先抠局部、怎样利用参考图。",
    "M2": "对着一组静物如实记录。看观察力与空间推理：前后遮挡、大小比例、物体之间的位置关系。",
    "M3": "由几块不完整的图形出发完成一幅画（残缺图形创造力范式）。看孩子如何把无意义的碎片赋予意义。",
    "M4": "把一件日常物品改造成另一种用途的东西。看变形与转化的想象力。",
    "M5": "把两个毫不相干的概念融合成一个新事物。看概念整合的能力，和 M4 的「改造一件」是不同的能力。",
    "M6": "用颜色改变同一个场景的情绪、天气或时间。让色彩的丰富与对比真正成为可观察的对象，而不是要求「多用颜色」。",
    "M7": "先用纯线条表现一个抽象概念，再把这些线发展成完整作品。看线条的组织与质感，以及从抽象到具象的转化。",
    "M8": "在给定的两条「世界规则」下画出那个世界的生活（初中版多一条可选的挑战规则）。最接近真实的自由创作，但由规则提供约束。",
    "M9": "开放的故事题。用来检验在受控任务里看到的行为模式，在真实创作中是否仍然存在。",
}

_WORD = {}
for tbl in (getattr(_m, "M4_BASE", []), getattr(_m, "M4_ENV", []), getattr(_m, "M4_GOAL", [])):
    for row in tbl:
        _WORD[row[0]] = row[-1]
for row in getattr(_m, "M5_PAIRS", []):
    _WORD[row[0]], _WORD[row[1]] = row[2], row[3]
for row in getattr(_m, "M7_CONCEPTS", []):
    _WORD[row[0]] = row[1]


TIER_ZH = {("simple", "full"): "两版共通", ("simple",): "小学版", ("full",): "初中版"}


def variant(q):
    c = q.get("condition") or {}
    bits = []
    t = tuple(q.get("tiers") or [])
    if t:
        bits.append(TIER_ZH.get(t, "/".join(t)))
    if q["family"] == "M4" and q.get("prompt_style") in STYLE_ZH:
        bits.append(STYLE_ZH[q["prompt_style"]])
    for k in ("mood_zh", "when_zh"):
        if c.get(k):
            bits.append(c[k])
    for k in ("base_object", "environment", "goal", "concept_a", "concept_b", "concept"):
        if c.get(k):
            bits.append(_WORD.get(c[k], c[k]))
    return " · ".join(bits)


def dims(q, key):
    return "、".join(ZH.get(k, k) for k in (q.get("rubric") or {}).get(key, []))


def build_doc():
    """清单的内容，和出口无关：[(家族标题, 研究目的, 主要考察, 其次, [行…])]。"""
    by_fam = OrderedDict()
    for q in QUESTS:
        if q.get("legacy"):
            continue          # v1 的 75 道留在库里给旧数据查 id，清单只列孩子现在看到的
        by_fam.setdefault(q["family"], []).append(q)
    fams = {f["id"]: f for f in families()}
    out = []
    for n, (fid, rows) in enumerate(by_fam.items(), 1):
        f = fams.get(fid, {})
        first = rows[0]
        table = [(f"{n}.{i}", q.get("title", ""), q.get("instruction") or "", q.get("hint") or "", variant(q))
                 for i, q in enumerate(rows, 1)]
        out.append((f"{n}. {f.get('name', '')}（{len(rows)} 题）", AIM.get(fid, f.get("research_goal", "")),
                    dims(first, "primary_dimensions"), dims(first, "secondary_dimensions"), table))
    return out


HEAD = ["编号", "标题", "孩子看到的指令", "提示", "变体"]
WIDTHS = [7, 15, 44, 20, 14]          # 列宽百分比，两个出口同一份
_N_LIVE = sum(1 for q in QUESTS if not q.get("legacy"))
INTRO = (f"共 {_N_LIVE} 个任务，分 10 个家族；每个家族小学版 5 道、初中版 5 道，其中 2 道两版共通。"
         "九个评价维度每幅画都评；每个家族另外标出它主要考察的维度。")


def build_html() -> str:
    e = html.escape
    CSS = """
    @page { size: A4; margin: 18mm 16mm; }
    body { font-family: 'Noto Sans CJK SC', 'Noto Serif CJK SC', 'PingFang SC', sans-serif; font-size: 10.5pt; line-height: 1.5; color: #222; }
    h1 { font-size: 20pt; margin: 0 0 4pt; } h2 { font-size: 14pt; margin: 22pt 0 6pt; page-break-after: avoid; }
    p { margin: 3pt 0; }
    table { border-collapse: collapse; width: 100%; margin: 6pt 0 4pt; table-layout: fixed; }
    th, td { border: 1px solid #999; padding: 4pt 5pt; vertical-align: top; text-align: left; font-size: 9.5pt; }
    th { background: #eee; }
    tr { page-break-inside: avoid; }
    """
    H = ["<html><head><meta charset='utf-8'><title>KidsArtQuest 任务清单</title><style>" + CSS + "</style></head><body>",
         "<h1>KidsArtQuest 任务清单</h1>", f"<p>{e(INTRO)}</p>",
         "<h2>九个维度</h2><table><colgroup><col width='18%'><col width='82%'></colgroup><tr><th>维度</th><th>看的是什么</th></tr>"]
    for d in DIMENSIONS:
        H.append(f"<tr><td><b>{e(d['zh'])}</b></td><td>{e(d['desc'])}</td></tr>")
    H.append("</table>")
    for title, aim, primary, sec, table in build_doc():
        H.append(f"<h2>{e(title)}</h2>")
        H.append(f"<p><b>研究目的</b>　{e(aim)}</p>")
        H.append(f"<p><b>主要考察</b>　{e(primary)}" + (f"　　<b>其次</b>　{e(sec)}" if sec else "") + "</p>")
        H.append("<table><colgroup>" + "".join(f"<col width='{w}%'>" for w in WIDTHS) + "</colgroup><tr>"
                 + "".join(f"<th>{h}</th>" for h in HEAD) + "</tr>")
        for row in table:
            cells = [e(c).replace("\n\n", "<br>").replace("\n", "<br>") for c in row]
            H.append("<tr>" + "".join(f"<td>{c}</td>" for c in cells) + "</tr>")
        H.append("</table>")
    H.append("</body></html>")
    return "\n".join(H)


def build_docx(path):
    """Word 版：python-docx 直接搭，列宽写死——LibreOffice 从 HTML 导入会把表格挤成一条。"""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    usable = sec.page_width - sec.left_margin - sec.right_margin
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.font.size = Pt(10.5)

    def table(head, rows, widths_pct, bold_first=False):
        t = doc.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        widths = [int(usable * w / 100) for w in widths_pct]
        for i, h in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5)
            shade = OxmlElement("w:shd"); shade.set(qn("w:val"), "clear"); shade.set(qn("w:fill"), "EEEEEE")
            c._tc.get_or_add_tcPr().append(shade)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                lines = str(val).split("\n")
                p = cells[i].paragraphs[0]
                for j, ln in enumerate(lines):
                    if not ln:
                        continue
                    if j and p.runs:
                        p = cells[i].add_paragraph()
                    r = p.add_run(ln); r.font.size = Pt(9.5)
                    if bold_first and i == 0:
                        r.bold = True
        # 列宽要写两处：每个单元格的 tcW（Word 看这个）和 tblGrid 的 gridCol（LibreOffice 看这个）
        for i, w in enumerate(widths):
            t.columns[i].width = w
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w
        layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(layout)
        # 行不跨页
        for row in t.rows[1:]:
            trPr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit"); trPr.append(cant)
        return t

    doc.add_heading("KidsArtQuest 任务清单", level=1)
    doc.add_paragraph(INTRO)
    doc.add_heading("九个维度", level=2)
    table(["维度", "看的是什么"], [(d["zh"], d["desc"]) for d in DIMENSIONS], [18, 82], bold_first=True)
    for title, aim, primary, sec_dims, rows in build_doc():
        doc.add_heading(title, level=2)
        p = doc.add_paragraph(); p.add_run("研究目的　").bold = True; p.add_run(aim)
        p = doc.add_paragraph(); p.add_run("主要考察　").bold = True; p.add_run(primary)
        if sec_dims:
            p.add_run("　　其次　").bold = True; p.add_run(sec_dims)
        table(HEAD, rows, WIDTHS)
    doc.save(path)


def main() -> int:
    out_docx = ROOT / "docs/任务清单.docx"
    out_pdf = ROOT / "docs/任务清单.pdf"
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "tasks.html"
        src.write_text(build_html(), encoding="utf-8")
        # PDF：Chrome 打印，CSS 说了算
        subprocess.run(["google-chrome", "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--print-to-pdf-no-header",
                        f"--print-to-pdf={out_pdf}", str(src)], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    build_docx(out_docx)
    print(f"wrote {out_docx.name} and {out_pdf.name}: {_N_LIVE} tasks")
    return 0


if __name__ == "__main__":
    sys.exit(main())
